"""Streamlit UI for the Sales Engineering Agent.

Landing page explains the tool. User enters a prospect name, uploads a discovery
transcript (.md, .txt, .pdf, or .docx), and clicks Analyse. The agent runs and
displays results in three tabs. Previously analysed prospects are listed in a
side column for quick switching.

Run with:
    uv run streamlit run app.py
"""

import io
import re
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent / ".env", override=True)

from sea.agent import run_agent

# ---------------------------------------------------------------------------
# Page config
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Sales Engineering Agent",
    page_icon="🔍",
    layout="wide",
)

# ---------------------------------------------------------------------------
# Session state
# ---------------------------------------------------------------------------

if "prospects" not in st.session_state:
    st.session_state.prospects = []      # list of {name, transcript_name, output}
if "active_idx" not in st.session_state:
    st.session_state.active_idx = None   # int index into prospects list
if "form_key" not in st.session_state:
    st.session_state.form_key = 0        # increment to reset name input + uploader

# ---------------------------------------------------------------------------
# Global CSS — navy palette, card styles, tab overrides
# ---------------------------------------------------------------------------

NAVY       = "#0f2244"
NAVY_MID   = "#1a3a6e"
NAVY_LIGHT = "#eef2f8"
ACCENT     = "#2563eb"
WHITE      = "#ffffff"

st.markdown(f"""
<style>
  /* ── Layout ── */
  .block-container {{ padding-top: 1.5rem !important; max-width: 1100px; }}
  [data-testid="stAppViewContainer"] {{ background: #f8f9fc; }}

  /* ── Hide Streamlit's auto-anchor icons on custom HTML headings ── */
  [data-testid="stHeadingAnchorLink"] {{ display: none !important; }}
  .hero h1 a, .step-card h4 a {{ display: none !important; }}

  /* ── Hero banner ── */
  .hero {{
    background: linear-gradient(135deg, {NAVY} 0%, {NAVY_MID} 100%);
    border-radius: 16px;
    padding: 3rem 2.75rem;
    margin-bottom: 2rem;
    color: {WHITE};
    position: relative;
    overflow: hidden;
  }}
  .hero-deco {{
    position: absolute;
    right: -50px;
    bottom: -60px;
    width: 340px;
    height: 340px;
    pointer-events: none;
    z-index: 1;
  }}
  .hero .hero-eyebrow,
  .hero h1,
  .hero p,
  .hero div {{ position: relative; z-index: 2; }}
  .hero-eyebrow {{
    font-size: 0.78rem;
    font-weight: 600;
    letter-spacing: 0.12em;
    text-transform: uppercase;
    color: rgba(255,255,255,0.55);
    margin-bottom: 0.6rem;
  }}
  .hero h1 {{
    color: {WHITE};
    font-size: 2.4rem;
    font-weight: 800;
    margin: 0 0 0.75rem;
    line-height: 1.15;
  }}
  .hero p {{
    color: rgba(255,255,255,0.78);
    font-size: 1.05rem;
    margin: 0;
    max-width: 620px;
    line-height: 1.6;
  }}

  /* ── Step cards ── */
  .steps-grid {{ display: flex; gap: 1rem; margin-bottom: 2rem; }}
  .step-card {{
    flex: 1;
    background: {WHITE};
    border-radius: 12px;
    border-top: 4px solid {NAVY};
    padding: 1.4rem 1.4rem 1.6rem;
    box-shadow: 0 1px 4px rgba(0,0,0,0.07);
  }}
  .step-card-top {{
    display: flex;
    justify-content: space-between;
    align-items: flex-start;
    margin-bottom: 0.85rem;
  }}
  .step-num {{
    display: inline-flex;
    align-items: center;
    justify-content: center;
    width: 26px; height: 26px;
    background: {NAVY};
    color: {WHITE};
    border-radius: 50%;
    font-size: 0.78rem;
    font-weight: 700;
    flex-shrink: 0;
  }}
  .step-icon {{
    width: 44px; height: 44px;
    background: {NAVY_LIGHT};
    border-radius: 10px;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    color: {NAVY};
    flex-shrink: 0;
  }}
  .step-card h4 {{
    margin: 0 0 0.3rem;
    font-size: 0.95rem;
    font-weight: 700;
    color: {NAVY};
  }}
  .step-card p {{
    margin: 0;
    font-size: 0.85rem;
    color: #4b5563;
    line-height: 1.5;
  }}

  /* ── Section headings ── */
  .section-label {{
    font-size: 0.72rem;
    font-weight: 700;
    letter-spacing: 0.1em;
    text-transform: uppercase;
    color: {NAVY};
    margin-bottom: 0.4rem;
  }}
  .section-bar {{
    height: 3px;
    background: {NAVY};
    border-radius: 2px;
    width: 40px;
    margin-bottom: 1rem;
  }}

  /* ── New Prospect card ── */
  .new-prospect-card {{
    background: {WHITE};
    border: 2px dashed {NAVY};
    border-radius: 14px;
    padding: 1.5rem 1.75rem;
    margin-bottom: 1.25rem;
  }}
  .new-prospect-card h3 {{
    color: {NAVY};
    font-size: 1.05rem;
    font-weight: 700;
    margin: 0 0 0.3rem;
  }}
  .new-prospect-card p {{
    color: #6b7280;
    font-size: 0.88rem;
    margin: 0;
  }}

  /* ── Your Prospects radio list ── */
  [data-testid="stRadio"] > div > label {{
    background: {WHITE};
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 0.55rem 0.9rem !important;
    margin-bottom: 0.4rem;
    cursor: pointer;
    font-size: 0.9rem;
    transition: border-color 0.15s;
  }}
  [data-testid="stRadio"] > div > label:hover {{
    border-color: {NAVY};
  }}

  /* ── File uploader ── */
  [data-testid="stFileUploader"] {{
    background: {WHITE};
    border: 2px dashed {NAVY};
    border-radius: 10px;
    padding: 0.5rem 1rem 1rem;
    margin-bottom: 0.75rem;
  }}
  [data-testid="stFileUploaderDropzone"] {{
    background: {NAVY_LIGHT} !important;
    border-color: #c7d4e8 !important;
    border-radius: 8px !important;
  }}

  /* ── Analyse transcript button ── */
  [data-testid="baseButton-primary"] {{
    background: {NAVY} !important;
    color: {WHITE} !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 600 !important;
    width: 100% !important;
    padding: 0.55rem 1.5rem !important;
  }}
  [data-testid="baseButton-primary"]:hover:not(:disabled) {{
    background: {NAVY_MID} !important;
  }}
  [data-testid="baseButton-primary"]:disabled {{
    opacity: 0.45 !important;
  }}

  /* ── Badges ── */
  .badge {{
    display: inline-block;
    padding: 4px 14px;
    border-radius: 20px;
    font-size: 0.82rem;
    font-weight: 600;
  }}
  .badge-navy   {{ background: {NAVY};    color: {WHITE}; }}
  .badge-blue   {{ background: {ACCENT};  color: {WHITE}; }}
  .badge-green  {{ background: #059669;   color: {WHITE}; }}
  .badge-orange {{ background: #d97706;   color: {WHITE}; }}
  .badge-gray   {{ background: #6b7280;   color: {WHITE}; }}

  /* ── Results section header ── */
  .results-header {{
    background: {NAVY};
    border-radius: 10px;
    padding: 0.9rem 1.4rem;
    color: {WHITE};
    font-weight: 700;
    font-size: 1rem;
    margin-bottom: 1.25rem;
    margin-top: 2rem;
  }}

  /* ── MEDDPICC letter badge inside expander ── */
  .medd-letter {{
    display: inline-block;
    width: 22px; height: 22px;
    background: {NAVY};
    color: {WHITE};
    border-radius: 4px;
    font-size: 0.75rem;
    font-weight: 800;
    text-align: center;
    line-height: 22px;
    margin-right: 0.4rem;
    flex-shrink: 0;
  }}

  /* ── Case study card ── */
  .cs-card {{
    background: {WHITE};
    border-radius: 12px;
    border: 1px solid #e2e8f0;
    border-top: 4px solid {NAVY};
    padding: 1.4rem;
    margin-bottom: 1rem;
    box-shadow: 0 1px 4px rgba(0,0,0,0.06);
  }}
  .cs-score {{
    float: right;
    background: {NAVY_LIGHT};
    color: {NAVY};
    font-size: 0.78rem;
    font-weight: 700;
    padding: 3px 10px;
    border-radius: 20px;
  }}

  /* ── Email card ── */
  .email-card {{
    background: {WHITE};
    border-radius: 12px;
    border: 1px solid #e2e8f0;
    padding: 2rem 2.25rem;
    box-shadow: 0 1px 4px rgba(0,0,0,0.06);
    white-space: pre-wrap;
    font-size: 0.95rem;
    line-height: 1.7;
    color: #1f2937;
  }}
  .email-subject {{
    font-size: 1.05rem;
    font-weight: 700;
    color: {NAVY};
    padding: 0.75rem 1.2rem;
    background: {NAVY_LIGHT};
    border-radius: 8px;
    margin-bottom: 1.25rem;
  }}

  /* ── Streamlit tab override ── */
  button[data-baseweb="tab"] {{
    font-weight: 600 !important;
    color: #4b5563 !important;
  }}
  button[data-baseweb="tab"][aria-selected="true"] {{
    color: {NAVY} !important;
    border-bottom-color: {NAVY} !important;
  }}

  /* ── Download buttons ── */
  .stDownloadButton > button {{
    background: {NAVY} !important;
    color: {WHITE} !important;
    border-radius: 8px !important;
    border: none !important;
    font-weight: 600 !important;
    padding: 0.5rem 1.25rem !important;
  }}
  .stDownloadButton > button:hover {{
    background: {NAVY_MID} !important;
  }}
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# Hero banner
# ---------------------------------------------------------------------------

st.markdown(f"""
<div class="hero">
  <div class="hero-eyebrow">Powered by Claude</div>
  <h1>Sales Engineering Agent</h1>
  <p>A Claude-powered agent built for <strong style="color:rgba(255,255,255,0.95)">Attributary</strong>,
  a fictional marketing data platform with two product lines: Core (data pipeline and warehouse connectors)
  and Measurement (marketing mix modelling and incrementality testing).</p>
  <p style="margin-top:0.75rem;">
  Upload a sales discovery call transcript and the agent runs three tools automatically: it extracts
  a structured MEDDPICC analysis, retrieves the most relevant customer case studies from the knowledge base,
  and drafts a tailored follow-up email grounded in both. No invented numbers. No hallucinated references.
  All three artefacts in under 60 seconds.</p>
  <div style="margin-top:1.5rem; padding-top:1.25rem; border-top:1px solid rgba(255,255,255,0.15);
              display:flex; gap:1.5rem; flex-wrap:wrap;">
    <span style="font-size:0.8rem; color:rgba(255,255,255,0.55);">
      Built with &nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Claude API</strong> &nbsp;·&nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Python</strong> &nbsp;·&nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Pydantic</strong> &nbsp;·&nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Streamlit</strong>
    </span>
    <span style="font-size:0.8rem; color:rgba(255,255,255,0.55);">
      Skills demonstrated &nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Tool use</strong> &nbsp;·&nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Structured outputs</strong> &nbsp;·&nbsp;
      <strong style="color:rgba(255,255,255,0.85);">RAG retrieval</strong> &nbsp;·&nbsp;
      <strong style="color:rgba(255,255,255,0.85);">Eval harness</strong>
    </span>
  </div>

  <!-- Decorative blob + icon -->
  <svg class="hero-deco" viewBox="0 0 360 360" fill="none" xmlns="http://www.w3.org/2000/svg">
    <path d="M185 28 C255 14 338 68 346 152 C354 236 298 318 214 330 C130 342 52 286 34 202 C16 118 58 38 122 24 C143 18 164 32 185 28 Z"
          fill="rgba(255,255,255,0.07)"/>
    <path d="M200 18 C272 4 355 58 361 144 C367 228 310 312 226 326 C142 340 62 282 44 196 C26 110 70 30 136 16 C158 10 178 22 200 18 Z"
          fill="none" stroke="rgba(147,197,253,0.28)" stroke-width="1.5"/>
    <circle cx="168" cy="162" r="58" stroke="rgba(255,255,255,0.16)" stroke-width="14" fill="none"/>
    <line x1="213" y1="207" x2="250" y2="244" stroke="rgba(255,255,255,0.16)" stroke-width="14" stroke-linecap="round"/>
  </svg>
</div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# What you get — three output cards
# ---------------------------------------------------------------------------

st.markdown(f"""
<div style="margin-bottom:0.75rem;">
  <span style="font-size:0.72rem; font-weight:700; letter-spacing:0.1em; text-transform:uppercase; color:{NAVY};">
    The outputs
  </span>
  <div style="height:3px; background:{NAVY}; border-radius:2px; margin-top:0.4rem; width:40px;"></div>
</div>
<div class="steps-grid">
  <div class="step-card">
    <div class="step-card-top">
      <div class="step-num">1</div>
      <div class="step-icon">
        <svg width="22" height="22" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
          <circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/>
        </svg>
      </div>
    </div>
    <h4>Structured discovery</h4>
    <p>Extracts MEDDPICC fields and classifies product-fit (Core vs Measurement) from spend and pain signals in the transcript.</p>
  </div>
  <div class="step-card">
    <div class="step-card-top">
      <div class="step-num">2</div>
      <div class="step-icon">
        <svg width="22" height="22" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
          <ellipse cx="12" cy="5" rx="9" ry="3"/>
          <path d="M21 12c0 1.66-4 3-9 3s-9-1.34-9-3"/>
          <path d="M3 5v14c0 1.66 4 3 9 3s9-1.34 9-3V5"/>
        </svg>
      </div>
    </div>
    <h4>Case study retrieval</h4>
    <p>Matches the discovery against a knowledge base of customer case studies, filtered and ranked by product line fit.</p>
  </div>
  <div class="step-card">
    <div class="step-card-top">
      <div class="step-num">3</div>
      <div class="step-icon">
        <svg width="22" height="22" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
          <rect x="2" y="4" width="20" height="16" rx="2"/>
          <path d="m22 7-8.97 5.7a1.94 1.94 0 0 1-2.06 0L2 7"/>
        </svg>
      </div>
    </div>
    <h4>Follow-up email</h4>
    <p>Drafts a tailored email grounded in the discovery and retrieved case studies. Every figure traces back to the transcript.</p>
  </div>
</div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------

def extract_text(file) -> str:
    name = file.name.lower()
    raw = file.read()

    if name.endswith(".pdf"):
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(raw))
        pages = [p.extract_text() for p in reader.pages if p.extract_text()]
        return "\n\n".join(pages)

    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(raw))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    return raw.decode("utf-8")


def exec_summary(content, max_len: int = 88) -> str:
    """Return a short one-liner from a field value for use as an expander label."""
    if not content:
        return ""
    if isinstance(content, list):
        text = str(content[0]) if content else ""
    else:
        text = str(content)
    # First sentence (stop at . ! ? or newline)
    m = re.match(r"([^.!?\n]+[.!?]?)", text.strip())
    first = m.group(1).strip() if m else text.strip()
    if len(first) > max_len:
        return first[:max_len].rstrip() + "..."
    return first


def format_meddpicc_writeup(d, fit_label: str) -> str:
    champ = d.champion
    if champ.status == "confirmed":
        champ_str = f"{champ.name} (confirmed)"
    elif champ.status == "forming":
        champ_str = f"{champ.name} (forming)"
    else:
        champ_str = "None identified"

    def block(title: str, content) -> str:
        if not content:
            return ""
        header = title.upper()
        sep = "-" * len(header)
        if isinstance(content, list):
            body = "\n".join(f"  * {item}" for item in content)
        else:
            body = f"  {content}"
        return f"{header}\n{sep}\n{body}\n"

    sections = [
        f"MEDDPICC DISCOVERY REPORT\n{'=' * 40}",
        f"Product Fit:  {fit_label}",
        f"Champion:     {champ_str}",
        "",
        block("Summary",          d.summary),
        block("Metrics",          d.metrics),
        block("Economic Buyer",   d.economic_buyer),
        block("Decision Criteria",d.decision_criteria),
        block("Decision Process", d.decision_process),
        block("Identify Pain",    d.identify_pain),
        block("Competition",      d.competition),
        block("Paper Process",    d.paper_process),
        block("Timeline",         d.timeline),
    ]
    return "\n".join(s for s in sections if s is not None)


def _sanitize(text: str) -> str:
    """Replace common unicode chars that fpdf2 core fonts cannot render."""
    return (text
        .replace("‘", "'").replace("’", "'")
        .replace("“", '"').replace("”", '"')
        .replace("–", "-").replace("—", "-")
        .replace("…", "...").replace(" ", " ")
        .replace("•", "*").replace("‐", "-"))


def _strip_md(text: str) -> str:
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    return text


def build_discovery_pdf(output, transcript_name: str, fit_label: str) -> bytes:
    from fpdf import FPDF

    d = output.discovery
    champ = d.champion
    if champ.status == "confirmed":
        champ_str = f"{champ.name} (confirmed)"
    elif champ.status == "forming":
        champ_str = f"{champ.name} (forming)"
    else:
        champ_str = "None identified"

    pdf = FPDF()
    pdf.set_margins(22, 22, 22)
    pdf.add_page()
    ew = pdf.epw

    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(15, 34, 68)
    pdf.cell(0, 12, "Discovery Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(107, 114, 128)
    pdf.cell(0, 5, _sanitize(transcript_name), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_draw_color(15, 34, 68)
    pdf.set_line_width(0.6)
    pdf.line(22, pdf.get_y(), pdf.w - 22, pdf.get_y())
    pdf.ln(5)

    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(35, 6, "Product Fit:", new_x="RIGHT", new_y="TOP")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, fit_label, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(35, 6, "Champion:", new_x="RIGHT", new_y="TOP")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, _sanitize(champ_str), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    def add_section(title: str, content):
        if not content:
            return
        pdf.set_fill_color(238, 242, 248)
        pdf.set_text_color(15, 34, 68)
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(0, 7, title.upper(), fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        pdf.set_text_color(31, 41, 55)
        pdf.set_font("Helvetica", "", 9)
        items = content if isinstance(content, list) else [content]
        for item in items:
            txt = _sanitize(str(item))
            prefix = "* " if isinstance(content, list) else ""
            pdf.set_x(27)
            pdf.multi_cell(ew - 5, 5, f"{prefix}{txt}")
        pdf.ln(3)

    add_section("Summary",           d.summary)
    add_section("Metrics",           d.metrics)
    add_section("Economic Buyer",    d.economic_buyer)
    add_section("Decision Criteria", d.decision_criteria)
    add_section("Decision Process",  d.decision_process)
    add_section("Identify Pain",     d.identify_pain)
    add_section("Competition",       d.competition)
    add_section("Paper Process",     d.paper_process)
    add_section("Timeline",          d.timeline)

    return bytes(pdf.output())


def build_case_studies_pdf(output, transcript_name: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_margins(22, 22, 22)
    pdf.add_page()
    ew = pdf.epw

    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(15, 34, 68)
    pdf.cell(0, 12, "Case Studies", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(107, 114, 128)
    pdf.cell(0, 5, _sanitize(transcript_name), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_draw_color(15, 34, 68)
    pdf.set_line_width(0.6)
    pdf.line(22, pdf.get_y(), pdf.w - 22, pdf.get_y())
    pdf.ln(6)

    for i, cs in enumerate(output.case_studies, 1):
        pdf.set_fill_color(15, 34, 68)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 9, f"{i}.  {_sanitize(cs.title)}", fill=True,
                 new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

        pdf.set_text_color(15, 34, 68)
        pdf.set_font("Helvetica", "B", 8)
        pdf.cell(0, 5, "WHY RETRIEVED", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(55, 65, 81)
        pdf.set_font("Helvetica", "", 8)
        for reason in cs.match_reasons:
            pdf.set_x(27)
            pdf.multi_cell(ew - 5, 4, f"* {_sanitize(reason)}")
        pdf.ln(3)

        pdf.set_text_color(15, 34, 68)
        pdf.set_font("Helvetica", "B", 8)
        pdf.cell(0, 5, "FULL CASE STUDY", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(31, 41, 55)
        pdf.set_font("Helvetica", "", 8)
        clean = _sanitize(_strip_md(cs.content))
        pdf.multi_cell(ew, 4, clean)
        pdf.ln(6)

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Agent dispatch (monkey-patched for progress indicators)
# ---------------------------------------------------------------------------

import sea.agent as _agent_mod

_original_make_dispatch = _agent_mod._make_dispatch

TOOL_LABELS = {
    "extract_structured_discovery": "Extracting MEDDPICC discovery...",
    "retrieve_case_studies":        "Retrieving case studies...",
    "draft_followup_email":         "Drafting follow-up email...",
}
TOOL_DONE = {
    "extract_structured_discovery": "Discovery extracted",
    "retrieve_case_studies":        "Case studies retrieved",
    "draft_followup_email":         "Email drafted",
}

# ---------------------------------------------------------------------------
# Two-column layout: "Your Prospects" | "New Prospect"
# ---------------------------------------------------------------------------

has_prospects = bool(st.session_state.prospects)

if has_prospects:
    col_your, col_new = st.columns([1, 2], gap="large")
else:
    col_your = None
    col_new = st.container()

# ── Your Prospects ──────────────────────────────────────────────────────────

if col_your is not None:
    with col_your:
        st.markdown(f"""
        <div class="section-label">Your Prospects</div>
        <div class="section-bar"></div>
        """, unsafe_allow_html=True)

        names = [p["name"] for p in st.session_state.prospects]
        default_idx = (
            st.session_state.active_idx
            if st.session_state.active_idx is not None
            else len(names) - 1
        )

        selected_idx = st.radio(
            "Select prospect",
            options=list(range(len(names))),
            format_func=lambda i: names[i],
            index=default_idx,
            label_visibility="collapsed",
        )
        st.session_state.active_idx = selected_idx

# ── New Prospect ─────────────────────────────────────────────────────────────

form_key = st.session_state.form_key

with col_new:
    st.markdown(f"""
    <div class="section-label">New Prospect</div>
    <div class="section-bar"></div>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="new-prospect-card">
      <h3>Upload a transcript</h3>
      <p>Enter the prospect or customer name, then upload the discovery call transcript.
         Accepts .md, .txt, .pdf, or .docx.</p>
    </div>
    """, unsafe_allow_html=True)

    prospect_name = st.text_input(
        "Prospect / Customer name",
        placeholder="e.g. Bellwether Coffee",
        key=f"pname_{form_key}",
    )

    uploaded = st.file_uploader(
        "Choose file",
        type=["md", "txt", "pdf", "docx"],
        label_visibility="collapsed",
        key=f"fup_{form_key}",
    )

    if uploaded:
        st.markdown(
            f'<p style="color:#6b7280; font-size:0.85rem; margin:0.25rem 0 0.75rem;">'
            f'<strong style="color:{NAVY};">{uploaded.name}</strong> &nbsp;·&nbsp; ready to analyse</p>',
            unsafe_allow_html=True,
        )

    can_run = bool(uploaded and prospect_name.strip())

    run_clicked = st.button(
        "Analyse transcript",
        disabled=not can_run,
        type="primary",
        key=f"run_{form_key}",
    )

    if run_clicked:
        transcript = extract_text(uploaded)

        if not transcript.strip():
            st.error("Could not extract text from the uploaded file. Please try a different format.")
            st.stop()

        _status_container = st.empty()

        def _make_verbose_dispatch(transcript_inner, state, client, model):
            inner = _original_make_dispatch(transcript_inner, state, client, model)
            def verbose_dispatch(tool_name, tool_input):
                with _status_container.status(
                    TOOL_LABELS.get(tool_name, tool_name), expanded=False
                ) as s:
                    result = inner(tool_name, tool_input)
                    s.update(label=TOOL_DONE.get(tool_name, tool_name), state="complete")
                return result
            return verbose_dispatch

        _agent_mod._make_dispatch = _make_verbose_dispatch

        with st.spinner("Running agent..."):
            try:
                output = run_agent(transcript)
            except Exception as e:
                st.error(f"Agent failed: {type(e).__name__}: {e}")
                st.stop()
            finally:
                _agent_mod._make_dispatch = _original_make_dispatch

        _status_container.empty()

        st.session_state.prospects.append({
            "name": prospect_name.strip(),
            "transcript_name": uploaded.name,
            "output": output,
        })
        st.session_state.active_idx = len(st.session_state.prospects) - 1
        st.session_state.form_key += 1  # reset name input + uploader on next render
        st.rerun()

# ---------------------------------------------------------------------------
# Results — shown for the selected prospect
# ---------------------------------------------------------------------------

if st.session_state.active_idx is not None and st.session_state.prospects:
    active          = st.session_state.prospects[st.session_state.active_idx]
    output          = active["output"]
    transcript_name = active["transcript_name"]

    st.markdown(
        f'<div class="results-header">Results &nbsp;·&nbsp; {active["name"]}</div>',
        unsafe_allow_html=True,
    )

    d = output.discovery

    fit_badge = {
        "core-only":        ("badge-blue",   "Core only"),
        "measurement-only": ("badge-orange",  "Measurement only"),
        "both":             ("badge-green",   "Core + Measurement"),
        "unclear":          ("badge-gray",    "Unclear"),
    }
    fit_cls, fit_label = fit_badge.get(d.product_fit, ("badge-gray", d.product_fit))

    tab_discovery, tab_cases, tab_email = st.tabs(
        ["Discovery", "Case Studies", "Follow-up Email"]
    )

    # ── Tab 1: Discovery ─────────────────────────────────────────────────────

    with tab_discovery:
        champ = d.champion
        if champ.status == "confirmed":
            champ_html = (
                f'<span class="badge badge-green">Champion confirmed</span>'
                f' &nbsp; {champ.name}'
            )
        elif champ.status == "forming":
            champ_html = (
                f'<span class="badge badge-orange">Champion forming</span>'
                f' &nbsp; {champ.name}'
            )
        else:
            champ_html = '<span class="badge badge-gray">No champion identified</span>'

        st.markdown(f"""
        <div style="display:flex; gap:1rem; align-items:center; flex-wrap:wrap; margin-bottom:1.5rem;">
          <span class="badge {fit_cls}">{fit_label}</span>
          {champ_html}
        </div>
        """, unsafe_allow_html=True)

        # ── MEDDPICC accordion in sequence ──────────────────────────────────

        # Champion content for the accordion
        if champ.status == "confirmed":
            champ_content: str | None = f"{champ.name} (confirmed)"
        elif champ.status == "forming":
            champ_content = f"{champ.name} (forming)"
        else:
            champ_content = None

        MEDDPICC_FIELDS = [
            ("M", "Metrics",           d.metrics),
            ("E", "Economic Buyer",    d.economic_buyer),
            ("D", "Decision Criteria", d.decision_criteria),
            ("D", "Decision Process",  d.decision_process),
            ("P", "Paper Process",     d.paper_process),
            ("I", "Identify Pain",     d.identify_pain),
            ("C", "Champion",          champ_content),
            ("C", "Competition",       d.competition),
        ]

        for letter, field_name, content in MEDDPICC_FIELDS:
            summary = exec_summary(content)
            if summary:
                label = f"**[{letter}]  {field_name}** &nbsp; {summary}"
            else:
                label = f"**[{letter}]  {field_name}**"

            with st.expander(label):
                if not content:
                    st.markdown(
                        "_Not identified in this transcript._",
                        unsafe_allow_html=False,
                    )
                elif isinstance(content, list):
                    for item in content:
                        st.markdown(f"- {item}")
                else:
                    st.markdown(str(content))

        # ── Full write-up ────────────────────────────────────────────────────
        st.markdown(
            f"<div style='margin-top:1.75rem; margin-bottom:0.4rem;'>"
            f"<span style='font-size:0.72rem; font-weight:700; letter-spacing:0.1em; "
            f"text-transform:uppercase; color:{NAVY};'>Full write-up</span>"
            f"<div style='height:2px; background:{NAVY}; border-radius:2px; "
            f"margin-top:0.3rem; width:40px;'></div>"
            f"</div>",
            unsafe_allow_html=True,
        )
        st.code(format_meddpicc_writeup(d, fit_label), language=None)

        st.download_button(
            label="Download Discovery PDF",
            data=build_discovery_pdf(output, transcript_name, fit_label),
            file_name="discovery.pdf",
            mime="application/pdf",
        )

    # ── Tab 2: Case Studies ───────────────────────────────────────────────────

    with tab_cases:
        if not output.case_studies:
            st.info("No case studies matched.")
        else:
            st.download_button(
                label="Download Case Studies PDF",
                data=build_case_studies_pdf(output, transcript_name),
                file_name="case-studies.pdf",
                mime="application/pdf",
            )
        for i, cs in enumerate(output.case_studies, 1):
            reasons_html = "".join(
                f'<li style="margin-bottom:0.25rem; font-size:0.88rem; color:#374151;">{r}</li>'
                for r in cs.match_reasons
            )
            st.markdown(f"""
            <div class="cs-card">
              <div class="cs-score">Score: {cs.score:.0f}</div>
              <div style="font-size:1rem; font-weight:700; color:{NAVY}; margin-bottom:0.5rem;">{i}. {cs.title}</div>
              <div style="font-size:0.75rem; font-weight:700; letter-spacing:0.08em; text-transform:uppercase;
                          color:{NAVY}; margin-bottom:0.35rem;">Why retrieved</div>
              <ul style="margin:0 0 0.5rem; padding-left:1.2rem;">{reasons_html}</ul>
            </div>
            """, unsafe_allow_html=True)
            with st.expander("View full case study"):
                st.markdown(cs.content)

    # ── Tab 3: Follow-up Email ────────────────────────────────────────────────

    with tab_email:
        st.markdown(
            f'<div class="email-subject">Subject: {output.email.subject}</div>',
            unsafe_allow_html=True,
        )
        body_html = output.email.body.replace("\n", "<br>")
        st.markdown(
            f'<div class="email-card">{body_html}</div>',
            unsafe_allow_html=True,
        )
        st.markdown("<div style='height:1rem'></div>", unsafe_allow_html=True)
        st.download_button(
            label="Download email",
            data=f"Subject: {output.email.subject}\n\n{output.email.body}",
            file_name="followup-email.txt",
            mime="text/plain",
        )
